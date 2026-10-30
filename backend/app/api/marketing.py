import io
import json
import re
import uuid
from calendar import monthrange
from datetime import date, datetime, timedelta
from typing import Any, Dict, List, Optional
from urllib.parse import quote

import pandas as pd
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.core.security import get_current_user, require_permission
from app.core.data_scope import apply_marketing_scope, require_all_scope
from app.models.app_config import AppConfig
from app.models.database import get_db
from app.models.employee import Employee
from app.models.marketing_detail import MarketingDetail
from app.utils.logger import log_operation

router = APIRouter(prefix="/api/marketing", tags=["营销统计"])

OVERALL_CONFIG_KEY = "marketing_overall_row_rules"
DEFAULT_OVERALL_CONFIG = {"rules": []}
KEY_PRODUCT_CONFIG_KEY = "marketing_key_products"
DEFAULT_KEY_PRODUCT_CONFIG = {"products": []}
UNCLASSIFIED_GROUP = "未分组"
UNASSIGNED_TEAM = "未分班组"
MATCH_TYPES = {"contains", "equals", "prefix", "suffix", "regex"}
WORK_GROUP_SUFFIXES = ("热线", "区域", "维系", "工单", "分公司")

OVERALL_SECTIONS = [
    {
        "section_no": 1,
        "section_name": "线上直办",
        "bg": "blue",
        "rows": ["线上直办-热线", "线上直办-区域", "线上直办-维系", "线上直办-工单", "线上直办-分公司"],
        "subtotal": True,
    },
    {
        "section_no": 2,
        "section_name": "营销工单",
        "bg": "green",
        "rows": ["营销工单-热线", "营销工单-区域", "营销工单-维系", "营销工单-工单", "营销工单-分公司"],
        "subtotal": True,
    },
    {
        "section_no": 3,
        "section_name": "降套维系",
        "bg": "yellow",
        "rows": ["降套维系"],
        "subtotal": False,
    },
]

OVERALL_ROW_LABELS = [row for section in OVERALL_SECTIONS for row in section["rows"]]
OVERALL_FIELDS = [
    "org_name",
    "channel",
    "pos_sequence",
    "develop_depart_name",
    "service_type",
    "data_type",
    "area_name",
    "tb_region_name",
    "tenant",
]
ONLINE_ROW_SUFFIXES = ("热线", "区域", "维系", "工单", "分公司")
ONLINE_ROW_LABELS = {f"线上直办-{suffix}" for suffix in ONLINE_ROW_SUFFIXES}
WORK_ORDER_ROW_LABELS = {f"营销工单-{suffix}" for suffix in ONLINE_ROW_SUFFIXES}
SOURCE_XSZB = "xszb"
SOURCE_FWGD = "fwgd"
SOURCE_LABELS = {SOURCE_XSZB: "线上直办", SOURCE_FWGD: "办理工单"}
PRODUCT_SOURCE_ALL = "all"
PRODUCT_SOURCE_LABELS = {
    SOURCE_XSZB: SOURCE_LABELS[SOURCE_XSZB],
    SOURCE_FWGD: SOURCE_LABELS[SOURCE_FWGD],
    PRODUCT_SOURCE_ALL: "合计",
}
OVERALL_COLUMNS = [
    {"key": "overlay", "label": "叠加包"},
    {"key": "package_change", "label": "套餐变更"},
    {"key": "broadband_new", "label": "宽带新装"},
    {"key": "mobile", "label": "移网"},
    {"key": "terminal", "label": "终端"},
]

FIELD_ALIASES = {
    "kf_id": ("KF_ID", "客服id", "客服ID", "工单受理工号"),
    "org_name": ("RESERVE6", "预留字段6"),
    "organization": ("ORG_NAME", "组织机构"),
    "tenant": ("建单租户",),
    "work_order_id": ("工单流水号",),
    "data_type": ("DATA_TYPE", "xszb/fwgd"),
    "pos_sequence": ("POS_SEQUENCE", "人员岗位性质"),
    "develop_depart_name": ("DEVELOP_DEPART_NAME", "发展渠道", "发展人渠道名称"),
    "develop_depart_id": ("DEVELOP_DEPART_ID", "发展渠道编码", "发展人渠道编码"),
    "develop_staff_name": ("DEVELOP_STAFF_NAME", "发展人", "发展人名称"),
    "develop_staff_id": ("DEVELOP_STAFF_ID", "发展人编码"),
    "tb_region_name": ("TB_REGION_NAME", "通报区域"),
    "channel": ("CHANNEL", "渠道名称", "受理渠道"),
    "area_name": ("AREA_NAME", "地市名称_cbss", "地市名称"),
    "net_type": ("NET_TYPE", "网别"),
    "trade_id": ("TRADE_ID", "CBSS订单号"),
    "finish_date": ("FINISH_DATE", "CBSS订单竣工时间"),
    "service_type": ("SERVICE_TYPE", "重点业务分类"),
    "package_type": ("PACKAGE_TYPE", "叠加包分类"),
    "valid_type": ("VALID_TYPE", "叠加包有效期"),
    "product_id": ("PRODUCT_ID", "产品ID"),
    "product_name": ("PRODUCT_NAME", "产品名称"),
    "lh_qh": ("LH_QH", "来话去话", "呼入方向"),
    "is_zero": ("IS_ZERO", "0元/非0元"),
}


class OverallRowRule(BaseModel):
    id: Optional[str] = None
    row_label: str
    field: str = "org_name"
    match_type: str = "contains"
    pattern: str = ""
    priority: int = 0
    enabled: bool = True


class OverallRulesIn(BaseModel):
    rules: List[OverallRowRule] = []


class KeyProductItem(BaseModel):
    id: Optional[str] = None
    business_name: str = ""
    product_id: str = ""
    product_name: str = ""
    sort: int = 0


class KeyProductsIn(BaseModel):
    products: List[KeyProductItem] = []


def _norm_label(value: Any) -> str:
    return str(value).replace("\ufeff", "").strip().lower()


def _text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.lower() in {"nan", "none", "nat", "--", "null"}:
        return ""
    return text


def _parse_finish_date(value: Any) -> Optional[date]:
    text = _text(value)
    if not text:
        return None
    for fmt in ("%Y%m%d%H%M%S", "%Y%m%d", "%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text[:len(datetime.now().strftime(fmt))], fmt).date()
        except Exception:
            pass
    parsed = pd.to_datetime(text, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.date()


def _read_excel(content: bytes):
    try:
        raw = pd.read_excel(io.BytesIO(content), header=None, dtype=object)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"无法解析Excel文件: {exc}")

    if raw.empty:
        raise HTTPException(status_code=400, detail="Excel文件为空")

    alias_lookup = {}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            alias_lookup[_norm_label(alias)] = field

    best_idx = 0
    best_hits = 0
    for idx in range(min(10, len(raw))):
        labels = [_norm_label(v) for v in raw.iloc[idx].tolist()]
        hits = sum(1 for label in labels if label in alias_lookup)
        if hits > best_hits:
            best_idx = idx
            best_hits = hits

    if best_hits < 3:
        raise HTTPException(status_code=400, detail="未识别到有效的营销明细表头")

    mapping = {}
    seen_fields = set()
    for col_idx, value in enumerate(raw.iloc[best_idx].tolist()):
        field = alias_lookup.get(_norm_label(value))
        if field and field not in seen_fields:
            mapping[field] = col_idx
            seen_fields.add(field)

    required = {"finish_date", "service_type"}
    missing = required - set(mapping)
    if missing:
        raise HTTPException(status_code=400, detail=f"Excel缺少必要列: {', '.join(sorted(missing))}")
    if "product_id" not in mapping and "product_name" not in mapping:
        raise HTTPException(status_code=400, detail="Excel缺少产品ID/产品名称列")

    return raw.iloc[best_idx + 1 :], mapping


def _cell(values: List[Any], mapping: Dict[str, int], field: str) -> str:
    idx = mapping.get(field)
    if idx is None or idx >= len(values):
        return ""
    return _text(values[idx])


def _rule_match(value: str, rule: Dict[str, Any]) -> bool:
    pattern = str(rule.get("pattern") or "")
    if not pattern:
        return False
    source = value or ""
    match_type = rule.get("match_type") or "contains"
    if match_type == "regex":
        try:
            return re.search(pattern, source, flags=re.IGNORECASE) is not None
        except re.error:
            return False
    source_cmp = source.lower()
    pattern_cmp = pattern.lower()
    if match_type == "equals":
        return source_cmp == pattern_cmp
    if match_type == "prefix":
        return source_cmp.startswith(pattern_cmp)
    if match_type == "suffix":
        return source_cmp.endswith(pattern_cmp)
    return pattern_cmp in source_cmp


def _get_config(db: Session, key: str, default: Dict[str, Any]) -> Dict[str, Any]:
    rec = db.query(AppConfig).filter(AppConfig.key == key).first()
    if not rec or not rec.value:
        return json.loads(json.dumps(default, ensure_ascii=False))
    try:
        data = json.loads(rec.value)
        if isinstance(data, dict):
            return data
    except Exception:
        pass
    return json.loads(json.dumps(default, ensure_ascii=False))


def _set_config(db: Session, key: str, data: Dict[str, Any]):
    rec = db.query(AppConfig).filter(AppConfig.key == key).first()
    serialized = json.dumps(data, ensure_ascii=False)
    if rec:
        rec.value = serialized
    else:
        db.add(AppConfig(key=key, value=serialized))


def _overall_config(db: Session) -> Dict[str, Any]:
    return _get_config(db, OVERALL_CONFIG_KEY, DEFAULT_OVERALL_CONFIG)


def _clean_rules(rules: List[Any]) -> List[Dict[str, Any]]:
    result = []
    for item in rules:
        data = item.model_dump() if hasattr(item, "model_dump") else dict(item)
        match_type = data.get("match_type") or "contains"
        if match_type not in MATCH_TYPES:
            raise HTTPException(status_code=400, detail=f"不支持的匹配方式: {match_type}")
        data["id"] = data.get("id") or uuid.uuid4().hex
        data["match_type"] = match_type
        data["priority"] = int(data.get("priority") or 0)
        data["enabled"] = bool(data.get("enabled", True))
        result.append(data)
    return result


def _key_product_config(db: Session) -> Dict[str, Any]:
    return _get_config(db, KEY_PRODUCT_CONFIG_KEY, DEFAULT_KEY_PRODUCT_CONFIG)


def _sorted_key_products(config: Dict[str, Any]) -> List[Dict[str, Any]]:
    products = config.get("products") or []
    return [
        product
        for _, product in sorted(
            enumerate(products),
            key=lambda pair: (int(pair[1].get("sort") or 0), pair[0]),
        )
    ]


def _norm_product_id(value: Any) -> str:
    return str(value or "").replace("\ufeff", "").strip().upper()


def _norm_product_name(value: Any) -> str:
    return str(value or "").replace("\ufeff", "").strip().casefold()


def _key_product_match_key(product: Dict[str, Any]) -> tuple:
    product_id = _norm_product_id(product.get("product_id"))
    if product_id:
        return ("id", product_id)
    return ("name", _norm_product_name(product.get("product_name")))


def _clean_key_products(products: List[Any]) -> List[Dict[str, Any]]:
    result = []
    seen = set()
    for index, item in enumerate(products):
        data = item.model_dump() if hasattr(item, "model_dump") else dict(item)
        business_name = str(data.get("business_name") or "").strip()
        product_id = str(data.get("product_id") or "").replace("\ufeff", "").strip()
        product_name = str(data.get("product_name") or "").replace("\ufeff", "").strip()
        if not business_name:
            raise HTTPException(status_code=400, detail="重点业务产品的业务名不能为空")
        if not product_id and not product_name:
            raise HTTPException(status_code=400, detail="重点业务产品需填写产品ID或产品名称")
        key = _key_product_match_key({"product_id": product_id, "product_name": product_name})
        if key in seen:
            raise HTTPException(status_code=400, detail=f"重点业务产品重复: {business_name}")
        seen.add(key)
        result.append({
            "id": data.get("id") or uuid.uuid4().hex,
            "business_name": business_name,
            "product_id": product_id,
            "product_name": product_name,
            "sort": index,
        })
    return result


def _normalize_product_source(value: Optional[str]) -> str:
    source = str(value or SOURCE_XSZB).strip().lower()
    if source not in (SOURCE_XSZB, SOURCE_FWGD, PRODUCT_SOURCE_ALL):
        raise HTTPException(status_code=400, detail="不支持的产品统计来源")
    return source


def _detail_source(detail: MarketingDetail) -> str:
    return SOURCE_FWGD if (detail.data_type or "").strip().lower() == SOURCE_FWGD else SOURCE_XSZB


def _source_condition(source: str):
    if source == SOURCE_FWGD:
        return MarketingDetail.data_type == SOURCE_FWGD
    return or_(
        MarketingDetail.data_type == SOURCE_XSZB,
        MarketingDetail.data_type.is_(None),
        MarketingDetail.data_type == "",
    )


def _detail_value(detail: MarketingDetail, field: str) -> str:
    if field == "org_name":
        org = str(getattr(detail, "org_name", "") or "").strip()
        if org:
            return org
        return str(getattr(detail, "organization", "") or "").strip()
    if field == "organization":
        org = str(getattr(detail, "organization", "") or "").strip()
        if org:
            return org
        return str(getattr(detail, "org_name", "") or "").strip()
    return str(getattr(detail, field, "") or "")


def _detail_values(detail: MarketingDetail) -> Dict[str, str]:
    return {field: _detail_value(detail, field) for field in OVERALL_FIELDS}


def _sorted_rules(config: Dict[str, Any]) -> List[Dict[str, Any]]:
    rules = [r for r in config.get("rules", []) if r.get("enabled", True)]
    return [rule for _, rule in sorted(enumerate(rules), key=lambda pair: (int(pair[1].get("priority") or 0), pair[0]))]


def _normalize_work_order_label(value: Any) -> Optional[str]:
    label = str(value or "").strip()
    if label in WORK_ORDER_ROW_LABELS:
        return label
    if label.startswith("线上直办-"):
        candidate = f"营销工单-{label.split('-', 1)[1]}"
        if candidate in WORK_ORDER_ROW_LABELS:
            return candidate
    return None


def _xszb_org_map(db: Session) -> Dict[str, str]:
    result: Dict[str, str] = {}
    details = db.query(MarketingDetail).filter(_source_condition(SOURCE_XSZB)).order_by(
        MarketingDetail.finish_date.desc(), MarketingDetail.id.desc()
    ).all()
    for detail in details:
        key = _norm_emp_no(detail.kf_id)
        if not key or key in result:
            continue
        org = (detail.organization or "").strip() or (detail.org_name or "").strip()
        if org:
            result[key] = org
    return result


def _classify_online_row(detail: MarketingDetail, config: Dict[str, Any], cache: Dict[str, str]) -> Optional[str]:
    values = _detail_values(detail)
    cache_key = SOURCE_XSZB + "\u0001" + "\u0001".join(values.get(field, "") for field in OVERALL_FIELDS)
    if cache_key in cache:
        return cache[cache_key] or None
    result = None
    for rule in _sorted_rules(config):
        field = rule.get("field") or "org_name"
        if field not in OVERALL_FIELDS:
            continue
        if _rule_match(values.get(field, ""), rule):
            label = str(rule.get("row_label") or "").strip()
            if label in ONLINE_ROW_LABELS or label == "降套维系":
                result = label
                break
    cache[cache_key] = result or ""
    return result


def _classify_work_order_row(
    detail: MarketingDetail,
    config: Dict[str, Any],
    org_map: Dict[str, str],
    cache: Dict[str, str],
) -> Optional[str]:
    values = _detail_values(detail)
    values["org_name"] = org_map.get(_norm_emp_no(detail.kf_id), "")
    cache_key = SOURCE_FWGD + "\u0001" + "\u0001".join(values.get(field, "") for field in OVERALL_FIELDS)
    if cache_key in cache:
        return cache[cache_key] or None

    rules = _sorted_rules(config)
    result = None
    for rule in rules:
        if (rule.get("field") or "org_name") != "org_name":
            continue
        if _rule_match(values.get("org_name", ""), rule):
            label = _normalize_work_order_label(rule.get("row_label"))
            if label:
                result = label
                break
    if not result:
        for rule in rules:
            if rule.get("field") != "tenant":
                continue
            if _rule_match(values.get("tenant", ""), rule):
                label = _normalize_work_order_label(rule.get("row_label"))
                if label:
                    result = label
                    break

    cache[cache_key] = result or ""
    return result


def _classify_overall_row(detail: MarketingDetail, config: Dict[str, Any], cache: Dict[str, str]) -> Optional[str]:
    if _detail_source(detail) == SOURCE_FWGD:
        return _classify_work_order_row(detail, config, {}, cache)
    return _classify_online_row(detail, config, cache)


def _resolve_range(year_month: Optional[str], start_date: Optional[str], end_date: Optional[str]):
    if year_month:
        start = datetime.strptime(f"{year_month}-01", "%Y-%m-%d").date()
        if year_month == datetime.now().strftime("%Y-%m"):
            end = datetime.now().date()
        else:
            next_month = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
            end = next_month - timedelta(days=1)
    elif start_date and end_date:
        start = datetime.strptime(start_date, "%Y-%m-%d").date()
        end = datetime.strptime(end_date, "%Y-%m-%d").date()
    elif start_date:
        start = datetime.strptime(start_date, "%Y-%m-%d").date()
        end = start
    elif end_date:
        end = datetime.strptime(end_date, "%Y-%m-%d").date()
        start = end
    else:
        now = datetime.now()
        _, last_day = monthrange(now.year, now.month)
        start = date(now.year, now.month, 1)
        end = date(now.year, now.month, last_day)
    return start, end


def _date_param(value: Optional[str]) -> Optional[date]:
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=400, detail="日期格式无效，应为 YYYY-MM-DD")


def _norm_emp_no(value: Any) -> str:
    return str(value or "").strip().upper()


def _emp_map(db: Session) -> Dict[str, Employee]:
    result = {}
    for emp in db.query(Employee).all():
        key = _norm_emp_no(emp.emp_no)
        if key:
            result.setdefault(key, emp)
    return result


def _team_leader_map(emps: Dict[str, Employee]) -> Dict[tuple, str]:
    leader_names: Dict[tuple, set] = {}
    for emp in emps.values():
        if emp.role == "组长" and emp.status == "在职" and emp.team and emp.name:
            key = ((emp.dept or "").strip(), emp.team)
            leader_names.setdefault(key, set()).add(emp.name)
    return {
        key: "、".join(sorted(names))
        for key, names in leader_names.items()
    }


def _extract_class(team: Optional[str]) -> str:
    match = re.match(r"^(.+?)[\d]+组$", (team or "").strip())
    return match.group(1) if match else (team or "").strip()


def _new_bucket():
    return {
        "total": 0,
        "inbound": 0,
        "outbound": 0,
        "overlay": 0,
        "overlay_day": 0,
        "overlay_month": 0,
        "overlay_long": 0,
        "overlay_undefined": 0,
        "package_change": 0,
        "orders": set(),
    }


def _business_key(service_type: str) -> Optional[str]:
    text = (service_type or "").strip()
    if text == "叠加包":
        return "overlay"
    if text in ("套餐适配", "套餐变更"):
        return "package_change"
    return None


def _valid_key(valid_type: str) -> str:
    text = (valid_type or "").strip()
    if text == "日包":
        return "overlay_day"
    if text == "月包":
        return "overlay_month"
    if text in ("长期", "长期包"):
        return "overlay_long"
    return "overlay_undefined"


def _accumulate(bucket: Dict[str, Any], detail: MarketingDetail):
    bucket["total"] += 1
    lh = (detail.lh_qh or "").strip()
    if "来" in lh:
        bucket["inbound"] += 1
    elif "去" in lh:
        bucket["outbound"] += 1
    biz = _business_key(detail.service_type or "")
    if biz == "overlay":
        bucket["overlay"] += 1
        bucket[_valid_key(detail.valid_type or "")] += 1
    elif biz == "package_change":
        bucket["package_change"] += 1
    if detail.trade_id:
        bucket["orders"].add(detail.trade_id)


def _finalize(bucket: Dict[str, Any]) -> Dict[str, Any]:
    data = {k: v for k, v in bucket.items() if k != "orders"}
    data["order_count"] = len(bucket.get("orders") or set())
    return data


def _add_to(bucket: Dict[str, Any], row: Dict[str, Any]):
    for key in (
        "total",
        "inbound",
        "outbound",
        "overlay",
        "overlay_day",
        "overlay_month",
        "overlay_long",
        "overlay_undefined",
        "package_change",
    ):
        bucket[key] = bucket.get(key, 0) + int(row.get(key) or 0)
    bucket.setdefault("orders", set())
    if row.get("_trade_id"):
        bucket["orders"].add(row["_trade_id"])


def _match_business_filter(detail: MarketingDetail, business_type: str) -> bool:
    if not business_type:
        return True
    biz = _business_key(detail.service_type or "")
    if business_type in ("叠加包", "overlay"):
        return biz == "overlay"
    if business_type in ("套餐变更", "套餐适配", "package_change"):
        return biz == "package_change"
    return True


def _match_valid_filter(detail: MarketingDetail, valid_type: str) -> bool:
    if not valid_type:
        return True
    if (detail.service_type or "").strip() != "叠加包":
        return False
    mapping = {
        "日包": "overlay_day",
        "月包": "overlay_month",
        "长期": "overlay_long",
        "长期包": "overlay_long",
        "未定义": "overlay_undefined",
        "day": "overlay_day",
        "month": "overlay_month",
        "long": "overlay_long",
        "undefined": "overlay_undefined",
    }
    return _valid_key(detail.valid_type or "") == mapping.get(valid_type)


def _compute_report(db: Session, start: date, end: date, filters: Dict[str, Any], current_user: dict) -> Dict[str, Any]:
    query = apply_marketing_scope(
        db.query(MarketingDetail).filter(
            MarketingDetail.finish_date >= start,
            MarketingDetail.finish_date <= end,
        ),
        db,
        current_user,
        "marketing_report",
    )
    data_type = filters.get("data_type") or ""
    if data_type in (SOURCE_XSZB, SOURCE_FWGD):
        query = query.filter(_source_condition(data_type))
    details = query.all()
    emps = _emp_map(db)
    team_leaders = _team_leader_map(emps)
    config = _overall_config(db)
    org_map = _xszb_org_map(db)
    group_cache: Dict[str, str] = {}

    prepared = []
    all_groups = set()
    all_teams = set()
    all_team_options = set()
    all_classes = set()
    all_class_options = set()
    all_orgs = set()
    all_products = set()
    for detail in details:
        emp = emps.get(_norm_emp_no(detail.kf_id))
        if _detail_source(detail) == SOURCE_FWGD:
            group = _classify_work_order_row(detail, config, org_map, group_cache) or UNCLASSIFIED_GROUP
            effective_org = org_map.get(_norm_emp_no(detail.kf_id), "")
        else:
            group = _classify_online_row(detail, config, group_cache) or UNCLASSIFIED_GROUP
            effective_org = _detail_value(detail, "org_name")
        prepared.append((detail, emp, group, effective_org))
        emp_dept = (emp.dept or "").strip() if emp else ""
        emp_team = emp.team if emp and emp.team else "未匹配班组"
        emp_class = _extract_class(emp_team)
        all_groups.add(group)
        all_teams.add(emp_team)
        all_team_options.add((emp_dept, emp_team))
        if emp_class:
            all_classes.add(emp_class)
            all_class_options.add((emp_dept, emp_class))
        if effective_org:
            all_orgs.add(effective_org)
        if detail.product_id or detail.product_name:
            all_products.add(detail.product_id or detail.product_name)

    group_name = filters.get("group_name") or ""
    dept = (filters.get("dept") or "").strip()
    team = filters.get("team") or ""
    class_name = filters.get("class_name") or ""
    name = filters.get("name") or ""
    emp_no = filters.get("emp_no") or ""
    org_name = filters.get("org_name") or ""
    business_type = filters.get("business_type") or ""
    valid_type = filters.get("valid_type") or ""
    lh_qh = filters.get("lh_qh") or ""
    product = filters.get("product") or ""

    filtered = []
    for detail, emp, group, effective_org in prepared:
        if group_name and group != group_name:
            continue
        if dept and (emp.dept or "").strip() != dept:
            continue
        if team:
            team_value = emp.team if emp and emp.team else "未匹配班组"
            if team_value != team:
                continue
        if class_name:
            team_value = emp.team if emp and emp.team else "未匹配班组"
            if _extract_class(team_value) != class_name:
                continue
        if name and not (emp and name.lower() in (emp.name or "").lower()):
            continue
        if emp_no and emp_no.lower() not in (detail.kf_id or "").lower():
            continue
        if org_name and org_name.lower() not in effective_org.lower():
            continue
        if lh_qh and (detail.lh_qh or "").strip() != lh_qh:
            continue
        if product:
            haystack = f"{detail.product_id or ''} {detail.product_name or ''}".lower()
            if product.lower() not in haystack:
                continue
        if not _match_business_filter(detail, business_type):
            continue
        if not _match_valid_filter(detail, valid_type):
            continue
        filtered.append((detail, emp, group, effective_org))

    employees: Dict[str, Any] = {}
    teams: Dict[str, Any] = {}
    groups: Dict[str, Any] = {}
    products: Dict[str, Any] = {}
    days: Dict[str, Any] = {}
    overall_bucket = _new_bucket()
    people = set()
    product_set = set()

    for detail, emp, group, effective_org in filtered:
        kf_id = detail.kf_id or ""
        emp_name = emp.name if emp else ""
        emp_team = emp.team if emp and emp.team else "未匹配班组"
        emp_dept = emp.dept if emp else ""
        if kf_id:
            people.add(kf_id)
        if detail.product_id or detail.product_name:
            product_set.add(detail.product_id or detail.product_name)

        emp_key = kf_id or "未匹配工号"
        bucket = employees.setdefault(emp_key, {
            **_new_bucket(),
            "kf_id": kf_id,
            "name": emp_name,
            "team": emp_team,
            "dept": emp_dept,
            "group_name": group,
            "org_name": effective_org,
        })
        if not bucket.get("tenant") and detail.tenant:
            bucket["tenant"] = detail.tenant
        _accumulate(bucket, detail)
        bucket["_trade_id"] = detail.trade_id

        team_key = (emp_dept, emp_team)
        team_bucket = teams.setdefault(team_key, {
            **_new_bucket(),
            "dept": emp_dept,
            "team": emp_team,
            "people": set(),
        })
        _accumulate(team_bucket, detail)
        team_bucket["people"].add(kf_id or emp_key)
        team_bucket["_trade_id"] = detail.trade_id

        group_bucket = groups.setdefault(group, {**_new_bucket(), "group_name": group, "people": set()})
        _accumulate(group_bucket, detail)
        group_bucket["people"].add(kf_id or emp_key)
        group_bucket["_trade_id"] = detail.trade_id

        product_key = detail.product_id or detail.product_name or "未匹配产品"
        product_bucket = products.setdefault(product_key, {
            **_new_bucket(),
            "product_id": detail.product_id or "",
            "product_name": detail.product_name or "",
        })
        _accumulate(product_bucket, detail)
        product_bucket["_trade_id"] = detail.trade_id

        day_key = detail.finish_date.isoformat() if detail.finish_date else ""
        day_bucket = days.setdefault(day_key, {**_new_bucket(), "date": day_key})
        _accumulate(day_bucket, detail)
        day_bucket["_trade_id"] = detail.trade_id

        _accumulate(overall_bucket, detail)

    def finalize_items(mapping, extra=None):
        result = []
        for key, bucket in mapping.items():
            item = _finalize(bucket)
            if extra:
                extra(item, bucket, key)
            result.append(item)
        result.sort(key=lambda item: (item.get("total", 0), item.get("order_count", 0)), reverse=True)
        return result

    def team_extra(item, bucket, key):
        item["dept"] = key[0]
        item["team"] = key[1]
        item["people"] = len(bucket.get("people") or set())
        item["leader"] = team_leaders.get(key, "")

    def group_extra(item, bucket, _key):
        item["people"] = len(bucket.get("people") or set())

    employee_items = finalize_items(employees)
    team_items = finalize_items(teams, team_extra)
    group_items = finalize_items(groups, group_extra)
    product_items = finalize_items(products)
    day_items = finalize_items(days)
    day_items.sort(key=lambda item: item.get("date") or "")

    stats = _finalize(overall_bucket)
    stats.update({
        "people_count": len(people),
        "product_count": len(product_set),
        "start_date": start.isoformat(),
        "end_date": end.isoformat(),
    })

    return {
        "stats": stats,
        "employees": employee_items,
        "teams": team_items,
        "groups": group_items,
        "products": product_items,
        "days": day_items,
        "options": {
            "groups": sorted(filter(None, all_groups)),
            "teams": sorted(filter(None, all_teams)),
            "team_options": [
                {"dept": dept_value, "team": team_value}
                for dept_value, team_value in sorted(all_team_options)
                if team_value
            ],
            "classes": sorted(filter(None, all_classes)),
            "class_options": [
                {"dept": dept_value, "class_name": class_value}
                for dept_value, class_value in sorted(all_class_options)
                if class_value
            ],
            "orgs": sorted(filter(None, all_orgs)),
            "products": sorted(filter(None, all_products)),
            "sources": [
                {"value": SOURCE_XSZB, "label": SOURCE_LABELS[SOURCE_XSZB]},
                {"value": SOURCE_FWGD, "label": SOURCE_LABELS[SOURCE_FWGD]},
            ],
        },
    }


def _overall_counts() -> Dict[str, int]:
    return {col["key"]: 0 for col in OVERALL_COLUMNS}


def _work_group_from_row_label(label: Optional[str]) -> Optional[str]:
    text = str(label or "").strip()
    if text.startswith("线上直办-"):
        text = text.split("-", 1)[1]
    elif text.startswith("营销工单-"):
        text = text.split("-", 1)[1]
    return text if text in WORK_GROUP_SUFFIXES else None


def _detail_matches_key_product(detail: MarketingDetail, product: Dict[str, Any]) -> bool:
    product_id = _norm_product_id(product.get("product_id"))
    if product_id:
        return _norm_product_id(detail.product_id) == product_id
    product_name = _norm_product_name(product.get("product_name"))
    return bool(product_name) and _norm_product_name(detail.product_name) == product_name


def _compute_key_product_table(
    db: Session,
    start: date,
    end: date,
    product_source: str,
    config: Dict[str, Any],
    org_map: Dict[str, str],
    cache: Dict[str, str],
    current_user: dict,
) -> Dict[str, Any]:
    products = _sorted_key_products(_key_product_config(db))
    columns = []
    for index, product in enumerate(products):
        product_id = str(product.get("product_id") or "").strip()
        product_name = str(product.get("product_name") or "").strip()
        columns.append({
            "key": f"key_product_{index}",
            "label": str(product.get("business_name") or product_name or product_id).strip(),
            "product_id": product_id,
            "product_name": product_name,
        })

    empty = {
        "key_product_columns": columns,
        "key_product_rows": [],
        "key_product_total_row": {"counts": {col["key"]: 0 for col in columns}, "total": 0},
        "key_product_source": product_source,
        "key_product_source_label": PRODUCT_SOURCE_LABELS.get(product_source, ""),
    }
    if not products:
        return empty

    query = apply_marketing_scope(
        db.query(MarketingDetail).filter(
            MarketingDetail.finish_date >= start,
            MarketingDetail.finish_date <= end,
        ),
        db,
        current_user,
        "marketing_overall",
    )
    if product_source in (SOURCE_XSZB, SOURCE_FWGD):
        query = query.filter(_source_condition(product_source))

    emps = _emp_map(db)
    team_leaders = _team_leader_map(emps)
    buckets: Dict[tuple, Dict[str, Any]] = {}
    for detail in query.all():
        if _detail_source(detail) == SOURCE_FWGD:
            row_label = _classify_work_order_row(detail, config, org_map, cache)
        else:
            row_label = _classify_online_row(detail, config, cache)
        work_group = _work_group_from_row_label(row_label)
        if not work_group:
            continue

        matched_keys = [
            column["key"]
            for column, product in zip(columns, products)
            if _detail_matches_key_product(detail, product)
        ]
        if not matched_keys:
            continue

        emp = emps.get(_norm_emp_no(detail.kf_id))
        team = (emp.team or "").strip() if emp and emp.team else UNASSIGNED_TEAM
        dept = (emp.dept or "").strip() if emp else ""
        bucket = buckets.setdefault(
            (work_group, dept, team),
            {"counts": {col["key"]: 0 for col in columns}, "total": 0},
        )
        for key in matched_keys:
            bucket["counts"][key] += 1
        bucket["total"] += len(matched_keys)

    group_order = {name: index for index, name in enumerate(WORK_GROUP_SUFFIXES)}
    sorted_items = sorted(
        buckets.items(),
        key=lambda pair: (
            group_order.get(pair[0][0], len(group_order)),
            pair[0][2] == UNASSIGNED_TEAM,
            -pair[1]["total"],
            pair[0][2],
            pair[0][1],
        ),
    )

    rows = []
    total_counts = {col["key"]: 0 for col in columns}
    total = 0
    for work_group in WORK_GROUP_SUFFIXES:
        group_items = [item for item in sorted_items if item[0][0] == work_group]
        has_subtotal = len(group_items) > 1
        group_counts = {col["key"]: 0 for col in columns}
        group_total = 0
        for index, ((_, dept, team), bucket) in enumerate(group_items):
            rows.append({
                "work_group": work_group,
                "dept": dept,
                "team": team,
                "leader": team_leaders.get((dept, team), ""),
                "rowspan": len(group_items) + (1 if has_subtotal else 0) if index == 0 else 0,
                "counts": dict(bucket["counts"]),
                "total": bucket["total"],
                "is_subtotal": False,
            })
            for key, value in bucket["counts"].items():
                total_counts[key] += value
                group_counts[key] += value
            total += bucket["total"]
            group_total += bucket["total"]

        if has_subtotal:
            rows.append({
                "work_group": work_group,
                "dept": "",
                "team": "小计",
                "rowspan": 0,
                "counts": group_counts,
                "total": group_total,
                "is_subtotal": True,
            })

    return {
        "key_product_columns": columns,
        "key_product_rows": rows,
        "key_product_total_row": {"counts": total_counts, "total": total},
        "key_product_source": product_source,
        "key_product_source_label": PRODUCT_SOURCE_LABELS.get(product_source, ""),
    }


def _source_date_ranges(
    db: Session,
    start: Optional[date] = None,
    end: Optional[date] = None,
    current_user: dict = None,
) -> Dict[str, Dict[str, Optional[str]]]:
    ranges = {}
    for source in (SOURCE_XSZB, SOURCE_FWGD):
        query = db.query(
            func.min(MarketingDetail.finish_date),
            func.max(MarketingDetail.finish_date),
        ).filter(_source_condition(source))
        if current_user:
            query = apply_marketing_scope(query, db, current_user, "marketing_overall")
        if start:
            query = query.filter(MarketingDetail.finish_date >= start)
        if end:
            query = query.filter(MarketingDetail.finish_date <= end)
        min_date, max_date = query.one()
        ranges[source] = {
            "start_date": min_date.isoformat() if min_date else None,
            "end_date": max_date.isoformat() if max_date else None,
        }
    return ranges


def _compute_overall(
    db: Session,
    start: date,
    end: date,
    product_source: str = SOURCE_XSZB,
    current_user: dict = None,
) -> Dict[str, Any]:
    config = _overall_config(db)
    cache: Dict[str, str] = {}
    org_map = _xszb_org_map(db)
    row_counts = {label: _overall_counts() for label in OVERALL_ROW_LABELS}
    unclassified_rows = 0
    unclassified_business_rows = 0

    online_query = db.query(MarketingDetail).filter(
        MarketingDetail.finish_date >= start,
        MarketingDetail.finish_date <= end,
        _source_condition(SOURCE_XSZB),
    )
    if current_user:
        online_query = apply_marketing_scope(online_query, db, current_user, "marketing_overall")
    online_details = online_query.all()
    for detail in online_details:
        biz = _business_key(detail.service_type or "")
        if biz not in ("overlay", "package_change"):
            unclassified_business_rows += 1
            continue
        row_label = _classify_online_row(detail, config, cache)
        if not row_label:
            unclassified_rows += 1
            continue
        if row_label not in ONLINE_ROW_LABELS:
            continue
        row_counts[row_label][biz] += 1

    work_order_query = db.query(MarketingDetail).filter(
        MarketingDetail.finish_date >= start,
        MarketingDetail.finish_date <= end,
        _source_condition(SOURCE_FWGD),
    )
    if current_user:
        work_order_query = apply_marketing_scope(work_order_query, db, current_user, "marketing_overall")
    work_order_details = work_order_query.all()
    for detail in work_order_details:
        biz = _business_key(detail.service_type or "")
        if biz not in ("overlay", "package_change"):
            unclassified_business_rows += 1
            continue
        row_label = _classify_work_order_row(detail, config, org_map, cache)
        if not row_label:
            unclassified_rows += 1
            continue
        row_counts[row_label][biz] += 1

    rows = []
    grand_total = 0
    for section in OVERALL_SECTIONS:
        item_rows = []
        for row_label in section["rows"]:
            counts = row_counts[row_label]
            total = sum(counts.values())
            item_rows.append({
                "section_no": section["section_no"],
                "section_name": section["section_name"],
                "bg": section["bg"],
                "label": row_label,
                "is_subtotal": False,
                "counts": counts,
                "total": total,
            })
        subtotal_counts = {
            col["key"]: sum(row["counts"][col["key"]] for row in item_rows)
            for col in OVERALL_COLUMNS
        }
        subtotal_total = sum(subtotal_counts.values())
        if section["subtotal"]:
            section_rows = item_rows + [{
                "section_no": section["section_no"],
                "section_name": section["section_name"],
                "bg": section["bg"],
                "label": "小计",
                "is_subtotal": True,
                "counts": subtotal_counts,
                "total": subtotal_total,
            }]
        else:
            section_rows = item_rows
        for idx, row in enumerate(section_rows):
            row["rowspan"] = (len(section_rows)) if idx == 0 else 0
            row["ratio"] = round(row["total"] / grand_total, 6) if grand_total else None
            rows.append(row)
        grand_total += subtotal_total

    for row in rows:
        row["ratio"] = round(row["total"] / grand_total, 6) if grand_total else None

    top_rows = []
    for section in OVERALL_SECTIONS:
        section_rows = [r for r in rows if r["section_no"] == section["section_no"]]
        if section["subtotal"]:
            top_rows.extend([r for r in section_rows if r["label"] == "小计"])
        else:
            top_rows.extend([r for r in section_rows if not r["is_subtotal"]])
    total_counts = {
        col["key"]: sum(row["counts"][col["key"]] for row in top_rows)
        for col in OVERALL_COLUMNS
    }

    total_row = {
        "label": "合计",
        "counts": total_counts,
        "total": sum(total_counts.values()),
        "ratio": 1 if total_counts else None,
    }
    key_product = _compute_key_product_table(
        db, start, end, product_source, config, org_map, cache, current_user
    )
    return {
        "stats": {
            "start_date": start.isoformat(),
            "end_date": end.isoformat(),
            "grand_total": grand_total,
            "unclassified_rows": unclassified_rows,
            "unclassified_business_rows": unclassified_business_rows,
        },
        "columns": OVERALL_COLUMNS,
        "rows": rows,
        "total_row": total_row,
        **key_product,
    }


@router.post("/details/import", response_model=dict)
def import_details(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing.upload")
    require_all_scope(current_user, "仅全部数据权限的账号可导入营销数据", "marketing")
    batch = uuid.uuid4().hex[:8]
    content = file.file.read()
    data, mapping = _read_excel(content)

    source = SOURCE_FWGD if ("tenant" in mapping or "work_order_id" in mapping) else SOURCE_XSZB
    records = []
    dates = set()
    seen = set()
    skipped = 0
    for pos in range(len(data)):
        values = data.iloc[pos].tolist()
        try:
            finish_date = _parse_finish_date(values[mapping["finish_date"]] if "finish_date" in mapping else None)
            if not finish_date:
                skipped += 1
                continue
            product_id = _cell(values, mapping, "product_id")
            product_name = _cell(values, mapping, "product_name")
            if not product_id and not product_name:
                skipped += 1
                continue
            org_name = _cell(values, mapping, "org_name")
            organization = _cell(values, mapping, "organization")
            if not org_name:
                org_name = organization
            if not organization:
                organization = org_name
            record = {
                "kf_id": _cell(values, mapping, "kf_id"),
                "org_name": org_name,
                "organization": organization,
                "tenant": _cell(values, mapping, "tenant") if source == SOURCE_FWGD else "",
                "work_order_id": _cell(values, mapping, "work_order_id") if source == SOURCE_FWGD else "",
                "data_type": source,
                "pos_sequence": _cell(values, mapping, "pos_sequence"),
                "develop_depart_name": _cell(values, mapping, "develop_depart_name"),
                "develop_depart_id": _cell(values, mapping, "develop_depart_id"),
                "develop_staff_name": _cell(values, mapping, "develop_staff_name"),
                "develop_staff_id": _cell(values, mapping, "develop_staff_id"),
                "tb_region_name": _cell(values, mapping, "tb_region_name"),
                "channel": _cell(values, mapping, "channel"),
                "area_name": _cell(values, mapping, "area_name"),
                "net_type": _cell(values, mapping, "net_type"),
                "trade_id": _cell(values, mapping, "trade_id"),
                "finish_date": finish_date,
                "service_type": _cell(values, mapping, "service_type"),
                "package_type": _cell(values, mapping, "package_type"),
                "valid_type": _cell(values, mapping, "valid_type"),
                "product_id": product_id,
                "product_name": product_name,
                "lh_qh": _cell(values, mapping, "lh_qh"),
                "is_zero": _cell(values, mapping, "is_zero"),
                "import_batch": batch,
            }
            key = (
                record["trade_id"],
                record["product_id"],
                record["product_name"],
                record["service_type"],
                record["valid_type"],
            )
            if record["trade_id"] and key in seen:
                skipped += 1
                continue
            if record["trade_id"]:
                seen.add(key)
            records.append(record)
            dates.add(finish_date)
        except Exception:
            skipped += 1

    if not records:
        raise HTTPException(status_code=400, detail="未解析到有效数据")

    db.query(MarketingDetail).filter(
        MarketingDetail.finish_date.in_(sorted(dates)),
        _source_condition(source),
    ).delete(synchronize_session=False)
    db.bulk_insert_mappings(MarketingDetail, records)
    db.commit()
    log_operation(
        db,
        current_user["id"],
        "import_marketing_details",
        "marketing",
        None,
        {"batch": batch, "count": len(records), "dates": [d.isoformat() for d in sorted(dates)], "source": source},
    )
    return {
        "count": len(records),
        "skipped": skipped,
        "batch": batch,
        "dates": [d.isoformat() for d in sorted(dates)],
        "data_type": source,
        "source": SOURCE_LABELS[source],
    }


@router.get("/details", response_model=dict)
def get_details(
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=200),
    import_batch: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    kf_id: Optional[str] = None,
    name: Optional[str] = None,
    org_name: Optional[str] = None,
    product: Optional[str] = None,
    service_type: Optional[str] = None,
    lh_qh: Optional[str] = None,
    data_type: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing.view")
    query = apply_marketing_scope(db.query(MarketingDetail), db, current_user, "marketing")
    if import_batch:
        query = query.filter(MarketingDetail.import_batch == import_batch)
    if data_type in (SOURCE_XSZB, SOURCE_FWGD):
        query = query.filter(_source_condition(data_type))
    start = _date_param(start_date)
    end = _date_param(end_date)
    if start:
        query = query.filter(MarketingDetail.finish_date >= start)
    if end:
        query = query.filter(MarketingDetail.finish_date <= end)
    if kf_id:
        query = query.filter(MarketingDetail.kf_id.ilike(f"%{kf_id}%"))
    if name:
        emp_nos = [e[0] for e in db.query(Employee.emp_no).filter(Employee.name.ilike(f"%{name}%")).all()]
        if not emp_nos:
            return {"items": [], "total": 0}
        query = query.filter(MarketingDetail.kf_id.in_(emp_nos))
    if org_name:
        org_map = _xszb_org_map(db)
        matched_kf = [key for key, value in org_map.items() if org_name.lower() in value.lower()]
        org_conditions = [
            MarketingDetail.org_name.ilike(f"%{org_name}%"),
            MarketingDetail.organization.ilike(f"%{org_name}%"),
        ]
        if matched_kf:
            org_conditions.append(MarketingDetail.kf_id.in_(matched_kf))
        query = query.filter(or_(*org_conditions))
    if product:
        query = query.filter(or_(
            MarketingDetail.product_id.ilike(f"%{product}%"),
            MarketingDetail.product_name.ilike(f"%{product}%"),
        ))
    if service_type:
        query = query.filter(MarketingDetail.service_type == service_type)
    if lh_qh:
        query = query.filter(MarketingDetail.lh_qh == lh_qh)

    total = query.count()
    items = query.order_by(
        MarketingDetail.finish_date.desc(), MarketingDetail.id.desc()
    ).offset((page - 1) * limit).limit(limit).all()

    emps = _emp_map(db)
    config = _overall_config(db)
    org_map = _xszb_org_map(db)
    cache: Dict[str, str] = {}
    result = []
    for detail in items:
        emp = emps.get(_norm_emp_no(detail.kf_id))
        source = _detail_source(detail)
        if source == SOURCE_FWGD:
            group = _classify_work_order_row(detail, config, org_map, cache) or UNCLASSIFIED_GROUP
            effective_org = org_map.get(_norm_emp_no(detail.kf_id), "")
        else:
            group = _classify_online_row(detail, config, cache) or UNCLASSIFIED_GROUP
            effective_org = detail.org_name or detail.organization or ""
        result.append({
            "id": detail.id,
            "kf_id": detail.kf_id,
            "name": emp.name if emp else "",
            "team": emp.team if emp else "",
            "dept": emp.dept if emp else "",
            "group_name": group,
            "org_name": effective_org,
            "organization": detail.organization,
            "tenant": detail.tenant,
            "work_order_id": detail.work_order_id,
            "data_type": source,
            "source": SOURCE_LABELS[source],
            "pos_sequence": detail.pos_sequence,
            "develop_depart_name": detail.develop_depart_name,
            "channel": detail.channel,
            "area_name": detail.area_name,
            "net_type": detail.net_type,
            "trade_id": detail.trade_id,
            "finish_date": detail.finish_date.isoformat() if detail.finish_date else None,
            "service_type": detail.service_type,
            "package_type": detail.package_type,
            "valid_type": detail.valid_type,
            "product_id": detail.product_id,
            "product_name": detail.product_name,
            "lh_qh": detail.lh_qh,
            "is_zero": detail.is_zero,
            "import_batch": detail.import_batch,
        })
    return {"items": result, "total": total}


@router.delete("/details/by-date", response_model=dict)
def delete_details_by_date(
    date: str = Query(...),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing.delete")
    require_all_scope(current_user, "仅全部数据权限的账号可删除营销数据", "marketing")
    target = _date_param(date)
    count = db.query(MarketingDetail).filter(MarketingDetail.finish_date == target).delete(synchronize_session=False)
    db.commit()
    log_operation(db, current_user["id"], "delete_marketing_details_by_date", "marketing", None, {"date": date, "count": count})
    return {"count": count}


@router.delete("/details/import/{batch}", response_model=dict)
def delete_details_batch(
    batch: str,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing.delete")
    require_all_scope(current_user, "仅全部数据权限的账号可删除营销数据", "marketing")
    count = db.query(MarketingDetail).filter(MarketingDetail.import_batch == batch).delete(synchronize_session=False)
    db.commit()
    log_operation(db, current_user["id"], "delete_marketing_batch", "marketing", None, {"batch": batch, "count": count})
    return {"count": count}


@router.delete("/details/{detail_id}", response_model=dict)
def delete_detail(
    detail_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing.delete")
    require_all_scope(current_user, "仅全部数据权限的账号可删除营销数据", "marketing")
    record = db.query(MarketingDetail).filter(MarketingDetail.id == detail_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="记录不存在")
    db.delete(record)
    db.commit()
    log_operation(db, current_user["id"], "delete_marketing_detail", "marketing", detail_id)
    return {"message": "删除成功"}


@router.get("/overall-row-rules", response_model=dict)
def get_overall_row_rules(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing_settings.view")
    return _overall_config(db)


@router.put("/overall-row-rules", response_model=dict)
def update_overall_row_rules(
    body: OverallRulesIn,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing_settings.edit")
    require_all_scope(current_user, "仅全部数据权限的账号可修改营销设置")
    cleaned = _clean_rules(body.rules)
    for rule in cleaned:
        if rule.get("row_label") not in OVERALL_ROW_LABELS:
            raise HTTPException(status_code=400, detail=f"不支持的行分类: {rule.get('row_label')}")
        if rule.get("field") not in OVERALL_FIELDS:
            raise HTTPException(status_code=400, detail=f"不支持的匹配字段: {rule.get('field')}")
    _set_config(db, OVERALL_CONFIG_KEY, {"rules": cleaned})
    db.commit()
    return _overall_config(db)


@router.get("/key-products", response_model=dict)
def get_key_products(
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing_settings.view")
    config = _key_product_config(db)
    return {"products": _sorted_key_products(config)}


@router.put("/key-products", response_model=dict)
def update_key_products(
    body: KeyProductsIn,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing_settings.edit")
    require_all_scope(current_user, "仅全部数据权限的账号可修改营销设置")
    cleaned = _clean_key_products(body.products)
    _set_config(db, KEY_PRODUCT_CONFIG_KEY, {"products": cleaned})
    db.commit()
    return {"products": cleaned}


@router.get("/product-options", response_model=dict)
def get_product_options(
    keyword: Optional[str] = None,
    limit: int = Query(200, ge=1, le=1000),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing_settings.view")
    query = apply_marketing_scope(db.query(
        MarketingDetail.product_id,
        MarketingDetail.product_name,
        func.count(MarketingDetail.id),
    ).filter(
        or_(
            MarketingDetail.product_id.isnot(None),
            MarketingDetail.product_name.isnot(None),
        )
    ), db, current_user, "marketing").group_by(
        MarketingDetail.product_id,
        MarketingDetail.product_name,
    ).order_by(func.count(MarketingDetail.id).desc())

    search = (keyword or "").strip().lower()
    items = []
    for product_id, product_name, count in query.all():
        pid = str(product_id or "").strip()
        pname = str(product_name or "").strip()
        if not pid and not pname:
            continue
        if search and search not in f"{pid} {pname}".lower():
            continue
        items.append({
            "product_id": pid,
            "product_name": pname,
            "count": int(count or 0),
        })
        if len(items) >= limit:
            break
    return {"items": items}


@router.get("/organizations", response_model=dict)
def get_organizations(
    limit: int = Query(200, ge=1, le=1000),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing_settings.view")
    config = _overall_config(db)
    cache: Dict[str, str] = {}
    org_map = _xszb_org_map(db)
    counter: Dict[tuple, int] = {}
    details = apply_marketing_scope(db.query(MarketingDetail), db, current_user, "marketing").all()
    for detail in details:
        source = _detail_source(detail)
        if source == SOURCE_FWGD:
            row_label = _classify_work_order_row(detail, config, org_map, cache) or UNCLASSIFIED_GROUP
            org_name = org_map.get(_norm_emp_no(detail.kf_id), "") or detail.tenant or ""
        else:
            row_label = _classify_online_row(detail, config, cache) or UNCLASSIFIED_GROUP
            org_name = _detail_value(detail, "org_name")
        key = (source, org_name, row_label)
        counter[key] = counter.get(key, 0) + 1
    items = [
        {
            "org_name": org_name,
            "count": count,
            "row_label": row_label,
            "data_type": source,
            "source": SOURCE_LABELS[source],
        }
        for (source, org_name, row_label), count in sorted(counter.items(), key=lambda item: item[1], reverse=True)[:limit]
    ]
    return {"items": items}


@router.get("/row-rule-preview", response_model=dict)
def get_row_rule_preview(
    limit: int = Query(100, ge=1, le=500),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing_settings.view")
    config = _overall_config(db)
    cache: Dict[str, str] = {}
    org_map = _xszb_org_map(db)
    details = apply_marketing_scope(db.query(MarketingDetail), db, current_user, "marketing").all()
    unclassified = 0
    counter: Dict[tuple, int] = {}
    for detail in details:
        if _business_key(detail.service_type or "") not in ("overlay", "package_change"):
            continue
        source = _detail_source(detail)
        if source == SOURCE_FWGD:
            row_label = _classify_work_order_row(detail, config, org_map, cache)
            values = _detail_values(detail)
            values["org_name"] = org_map.get(_norm_emp_no(detail.kf_id), "")
        else:
            row_label = _classify_online_row(detail, config, cache)
            values = _detail_values(detail)
        if row_label:
            continue
        unclassified += 1
        for field in OVERALL_FIELDS:
            value = values.get(field, "")
            if value:
                key = (field, value)
                counter[key] = counter.get(key, 0) + 1
    samples = [
        {"field": field, "value": value, "count": count}
        for (field, value), count in sorted(counter.items(), key=lambda item: item[1], reverse=True)[:limit]
    ]
    return {"unclassified": unclassified, "samples": samples}


@router.get("/report", response_model=dict)
def get_report(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    year_month: Optional[str] = None,
    group_name: Optional[str] = None,
    dept: Optional[str] = None,
    team: Optional[str] = None,
    class_name: Optional[str] = None,
    name: Optional[str] = None,
    emp_no: Optional[str] = None,
    org_name: Optional[str] = None,
    business_type: Optional[str] = None,
    valid_type: Optional[str] = None,
    lh_qh: Optional[str] = None,
    product: Optional[str] = None,
    data_type: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing_report.view")
    start, end = _resolve_range(year_month, start_date, end_date)
    return _compute_report(db, start, end, {
        "group_name": group_name,
        "dept": dept,
        "team": team,
        "class_name": class_name,
        "name": name,
        "emp_no": emp_no,
        "org_name": org_name,
        "business_type": business_type,
        "valid_type": valid_type,
        "lh_qh": lh_qh,
        "product": product,
        "data_type": data_type,
    }, current_user)


@router.get("/report/export")
def export_report(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    year_month: Optional[str] = None,
    dimension: str = Query("employee"),
    group_name: Optional[str] = None,
    dept: Optional[str] = None,
    team: Optional[str] = None,
    class_name: Optional[str] = None,
    name: Optional[str] = None,
    emp_no: Optional[str] = None,
    org_name: Optional[str] = None,
    business_type: Optional[str] = None,
    valid_type: Optional[str] = None,
    lh_qh: Optional[str] = None,
    product: Optional[str] = None,
    data_type: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing_report.export")
    start, end = _resolve_range(year_month, start_date, end_date)
    data = _compute_report(db, start, end, {
        "group_name": group_name,
        "dept": dept,
        "team": team,
        "class_name": class_name,
        "name": name,
        "emp_no": emp_no,
        "org_name": org_name,
        "business_type": business_type,
        "valid_type": valid_type,
        "lh_qh": lh_qh,
        "product": product,
        "data_type": data_type,
    }, current_user)
    output = io.StringIO()
    import csv

    writer = csv.writer(output)
    common = ["产品办理量", "订单数", "来话办理量", "去话办理量", "叠加包", "日包", "月包", "长期包", "未定义", "套餐变更"]
    if dimension == "team":
        writer.writerow(["部门", "班组", "组长", "人数"] + common)
        for item in data["teams"]:
            writer.writerow([item.get("dept", ""), item.get("team"), item.get("leader", ""), item.get("people", 0)] + [item.get(k, 0) for k in ("total", "order_count", "inbound", "outbound", "overlay", "overlay_day", "overlay_month", "overlay_long", "overlay_undefined", "package_change")])
    elif dimension == "group":
        writer.writerow(["行分类", "人数"] + common)
        for item in data["groups"]:
            writer.writerow([item.get("group_name"), item.get("people", 0)] + [item.get(k, 0) for k in ("total", "order_count", "inbound", "outbound", "overlay", "overlay_day", "overlay_month", "overlay_long", "overlay_undefined", "package_change")])
    elif dimension == "product":
        writer.writerow(["产品ID", "产品名称"] + common)
        for item in data["products"]:
            writer.writerow([item.get("product_id"), item.get("product_name")] + [item.get(k, 0) for k in ("total", "order_count", "inbound", "outbound", "overlay", "overlay_day", "overlay_month", "overlay_long", "overlay_undefined", "package_change")])
    elif dimension == "day":
        writer.writerow(["日期"] + common)
        for item in data["days"]:
            writer.writerow([item.get("date")] + [item.get(k, 0) for k in ("total", "order_count", "inbound", "outbound", "overlay", "overlay_day", "overlay_month", "overlay_long", "overlay_undefined", "package_change")])
    else:
        writer.writerow(["工号", "姓名", "班组", "部门", "行分类"] + common)
        for item in data["employees"]:
            writer.writerow([item.get("kf_id"), item.get("name"), item.get("team"), item.get("dept"), item.get("group_name")] + [item.get(k, 0) for k in ("total", "order_count", "inbound", "outbound", "overlay", "overlay_day", "overlay_month", "overlay_long", "overlay_undefined", "package_change")])

    filename = f"marketing_report_{start}_{end}_{dimension}.csv"
    log_operation(db, current_user["id"], "export_marketing_report", "marketing", None, {"start_date": start.isoformat(), "end_date": end.isoformat(), "dimension": dimension})
    return StreamingResponse(
        iter(["\ufeff" + output.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/overall-report", response_model=dict)
def get_overall_report(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    year_month: Optional[str] = None,
    product_source: Optional[str] = SOURCE_XSZB,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing_overall.view")
    start, end = _resolve_range(year_month, start_date, end_date)
    source = _normalize_product_source(product_source)
    data = _compute_overall(db, start, end, source, current_user)
    data["data_ranges"] = {
        "database": _source_date_ranges(db, current_user=current_user),
        "query": _source_date_ranges(db, start, end, current_user=current_user),
    }
    return data


def _excel_border():
    side = Side(style="thin", color="FF000000")
    return Border(left=side, right=side, top=side, bottom=side)


def _fill(color: str):
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _style_cell(cell, fill_color=None, bold=False, font_color=None):
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _excel_border()
    if fill_color:
        cell.fill = _fill(fill_color)
    if font_color or bold:
        cell.font = Font(bold=bold, color=font_color or "FF000000")


@router.get("/overall-report/export")
def export_overall_report(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    year_month: Optional[str] = None,
    product_source: Optional[str] = SOURCE_XSZB,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    require_permission(current_user, "marketing_overall.export")
    start, end = _resolve_range(year_month, start_date, end_date)
    source = _normalize_product_source(product_source)
    data = _compute_overall(db, start, end, source, current_user)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "总体营销"

    headers = ["序号", "五大类业务", "叠加包", "套餐变更", "宽带新装", "移网", "终端", "合计", "占比"]
    for col_idx, label in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=label)
        font_color = "FFFF0000" if 3 <= col_idx <= 7 else "FF000000"
        _style_cell(cell, fill_color="FFFFFFFF", bold=True, font_color=font_color)

    colors = {
        "blue": "FFDDEBF7",
        "pink": "FFFCE4EC",
        "green": "FFE2EFDA",
        "yellow": "FFFFFF00",
    }

    row_number = 2
    section_first_rows = {}
    merge_ranges = []
    total_row_number = 2 + len(data["rows"])

    for section in OVERALL_SECTIONS:
        section_rows = [row for row in data["rows"] if row["section_no"] == section["section_no"]]
        section_first_rows[section["section_no"]] = row_number
        for row in section_rows:
            fill_color = colors.get(row["bg"])
            sheet.cell(row=row_number, column=1, value=row["section_no"] if row["rowspan"] else None)
            sheet.cell(row=row_number, column=2, value=row["label"])
            for offset, column in enumerate(OVERALL_COLUMNS):
                sheet.cell(row=row_number, column=3 + offset, value=row["counts"][column["key"]])
            sheet.cell(row=row_number, column=8, value=f"=SUM(C{row_number}:G{row_number})")
            sheet.cell(row=row_number, column=9, value=f"=IFERROR(H{row_number}/$H${total_row_number},0)")
            for col_idx in range(1, 10):
                cell = sheet.cell(row=row_number, column=col_idx)
                _style_cell(cell, fill_color=fill_color, bold=bool(row.get("is_subtotal")))
                if col_idx == 9 and data["stats"]["grand_total"] == 0:
                    cell.font = Font(color="FF00B050", bold=True)
                    cell.fill = _fill("FFC6EFCE")
            if row["rowspan"] > 1:
                merge_ranges.append((row_number, row_number + row["rowspan"] - 1))
            row_number += 1

    for start_row, end_row in merge_ranges:
        sheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)

    total_row_number = row_number
    for col_idx, value in enumerate([None, None, 0, 0, 0, 0, 0, 0, 1], start=1):
        cell = sheet.cell(row=total_row_number, column=col_idx)
        if col_idx == 1:
            cell.value = "合计"
        elif col_idx >= 3 and col_idx <= 7:
            col_letter = chr(ord("A") + col_idx - 1)
            refs = []
            for section in OVERALL_SECTIONS:
                if section["subtotal"]:
                    refs.append(f"{col_letter}{section_first_rows[section['section_no']] + len(section['rows'])}")
                else:
                    refs.append(f"{col_letter}{section_first_rows[section['section_no']]}")
            cell.value = "=" + "+".join(refs)
        elif col_idx == 8:
            cell.value = f"=SUM(C{total_row_number}:G{total_row_number})"
        elif col_idx == 9:
            cell.value = 1
        _style_cell(cell, fill_color="FFFFFFFF", bold=True)

    sheet.merge_cells(start_row=total_row_number, start_column=1, end_row=total_row_number, end_column=2)

    key_columns = data.get("key_product_columns") or []
    if key_columns:
        row_number = total_row_number + 2
        last_column = 2 + len(key_columns) + 1
        title = f"重点业务产品统计（{data.get('key_product_source_label') or PRODUCT_SOURCE_LABELS.get(source, '')}）"
        sheet.cell(row=row_number, column=1, value=title)
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=last_column)
        for col_idx in range(1, last_column + 1):
            _style_cell(sheet.cell(row=row_number, column=col_idx), fill_color="FFDDEBF7", bold=True)
        row_number += 1

        key_headers = ["工作组", "班组"] + [column["label"] for column in key_columns] + ["合计"]
        for col_idx, label in enumerate(key_headers, start=1):
            cell = sheet.cell(row=row_number, column=col_idx, value=label)
            _style_cell(cell, fill_color="FFFFFFFF", bold=True, font_color="FFFF0000" if 3 <= col_idx <= last_column - 1 else None)
        row_number += 1

        group_merge_ranges = []
        for item in data.get("key_product_rows") or []:
            team_label = item["team"]
            if item.get("leader"):
                team_label = f"{team_label}（{item['leader']}）"
            sheet.cell(row=row_number, column=1, value=item["work_group"] if item.get("rowspan") else None)
            sheet.cell(row=row_number, column=2, value=team_label)
            for offset, column in enumerate(key_columns):
                sheet.cell(row=row_number, column=3 + offset, value=item["counts"].get(column["key"], 0))
            sheet.cell(
                row=row_number,
                column=last_column,
                value=f"=SUM(C{row_number}:{get_column_letter(last_column - 1)}{row_number})",
            )
            for col_idx in range(1, last_column + 1):
                if item.get("is_subtotal"):
                    _style_cell(sheet.cell(row=row_number, column=col_idx), fill_color="FFF2F2F2", bold=True)
                else:
                    _style_cell(sheet.cell(row=row_number, column=col_idx))
            if item.get("rowspan", 0) > 1:
                group_merge_ranges.append((row_number, row_number + item["rowspan"] - 1))
            row_number += 1

        for start_row, end_row in group_merge_ranges:
            sheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)

        key_total_row = row_number
        sheet.cell(row=key_total_row, column=1, value="合计")
        sheet.merge_cells(start_row=key_total_row, start_column=1, end_row=key_total_row, end_column=2)
        key_total = data.get("key_product_total_row") or {"counts": {}, "total": 0}
        for offset, column in enumerate(key_columns):
            sheet.cell(row=key_total_row, column=3 + offset, value=key_total["counts"].get(column["key"], 0))
        sheet.cell(
            row=key_total_row,
            column=last_column,
            value=f"=SUM(C{key_total_row}:{get_column_letter(last_column - 1)}{key_total_row})",
        )
        for col_idx in range(1, last_column + 1):
            _style_cell(sheet.cell(row=key_total_row, column=col_idx), fill_color="FFFFFFFF", bold=True)

    widths = {"A": 7, "B": 24, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    filename = f"总体营销_{start}_{end}.xlsx"
    encoded = quote(filename)
    log_operation(db, current_user["id"], "export_marketing_overall", "marketing", None, {"start_date": start.isoformat(), "end_date": end.isoformat(), "product_source": source})
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )
