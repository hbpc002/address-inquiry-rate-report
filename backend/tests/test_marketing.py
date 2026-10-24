import io
import json
import os
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

os.environ.setdefault("DATABASE_URL", "postgresql://postgres:admin123%40kf@localhost:5432/schedule_test")

from fastapi.testclient import TestClient
from sqlalchemy import text

import app.models  # noqa: F401
from app.core.security import get_current_user, get_password_hash
from app.main import app
from app.models.app_config import AppConfig
from app.models.database import Base, SessionLocal, engine, init_db
from app.models.employee import Employee
from app.models.marketing_detail import MarketingDetail
from app.models.operation_log import OperationLog
from app.models.user import User

ADMIN = {"id": 1, "username": "marketing_admin", "role": "admin", "is_system": True, "permissions": "{}"}


def _override_user():
    return ADMIN


app.dependency_overrides[get_current_user] = _override_user
client = TestClient(app)

EN_HEADERS = [
    "KF_ID", "RESERVE6", "DATA_TYPE", "POS_SEQUENCE", "DEVELOP_DEPART_NAME",
    "CHANNEL", "AREA_NAME", "TRADE_ID", "FINISH_DATE", "SERVICE_TYPE",
    "PACKAGE_TYPE", "VALID_TYPE", "PRODUCT_ID", "PRODUCT_NAME", "LH_QH",
]

CN_HEADERS = [
    "客服id", "预留字段6", "重点业务分类", "叠加包分类", "叠加包有效期",
    "CBSS订单竣工时间", "产品ID", "产品名称", "来话去话", "CBSS订单号",
    "受理渠道", "人员岗位性质", "发展渠道",
]

FWGD_HEADERS = [
    "工单受理工号", "建单租户", "工单流水号", "CBSS订单号", "CBSS订单竣工时间",
    "重点业务分类", "叠加包分类", "叠加包有效期", "产品ID", "产品名称",
    "呼入方向", "渠道名称", "人员岗位性质", "发展人渠道名称", "通报区域", "xszb/fwgd",
]

ORG_HEADERS = [
    "KF_ID", "组织机构", "CBSS订单竣工时间", "重点业务分类", "叠加包有效期",
    "产品ID", "产品名称", "CBSS订单号",
]


def setup_module():
    Base.metadata.drop_all(bind=engine)
    init_db()
    db = SessionLocal()
    try:
        if not db.query(User).filter(User.id == 1).first():
            db.add(User(id=1, username="marketing_admin", password_hash=get_password_hash("x"), role="admin", is_active=True))
            db.commit()
    finally:
        db.close()


def teardown_module():
    app.dependency_overrides.clear()


def _clean():
    db = SessionLocal()
    try:
        db.query(MarketingDetail).delete()
        db.query(Employee).delete()
        db.query(AppConfig).filter(AppConfig.key.in_(["marketing_overall_row_rules", "marketing_key_products"])).delete(synchronize_session=False)
        db.query(OperationLog).delete()
        db.commit()
    finally:
        db.close()


def _seed_employees():
    db = SessionLocal()
    try:
        db.add_all([
            Employee(emp_no="KF001", name="张三", team="热线一组", dept="客服部", status="在职"),
            Employee(emp_no="KF002", name="李四", team="区域二组", dept="客服部", status="在职"),
        ])
        db.commit()
    finally:
        db.close()


def _seed_leader(emp_no, name, team, status="在职"):
    db = SessionLocal()
    try:
        db.add(Employee(emp_no=emp_no, name=name, team=team, dept="客服部", role="组长", status=status))
        db.commit()
    finally:
        db.close()


def _make_xlsx(headers, rows, pre_rows=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in pre_rows or []:
        ws.append(row)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _en_row(kf="KF001", org="广西分公司>>省中心>>客户服务营销中心>>热线运营组",
            service="叠加包", valid="月包", finish="20261001120000", product_id="P1",
            product_name="测试产品", lh="来话", trade="T1", channel="热线人工",
            data_type="xszb", pos="一线坐席-10010", depart="10010", area="南宁"):
    return [kf, org, data_type, pos, depart, channel, area, trade, finish,
            service, "流量", valid, product_id, product_name, lh]


def _cn_row(kf="KF001", org="广西分公司>>省中心>>客户服务营销中心>>热线运营组",
            service="叠加包", valid="月包", finish="20261002120000", product_id="P1",
            product_name="测试产品", lh="来话", trade="T1", channel="热线人工",
            pos="一线坐席-10010", depart="10010"):
    return [kf, org, service, "流量", valid, finish, product_id, product_name,
            lh, trade, channel, pos, depart]


def _fwgd_row(kf="KF001", tenant="广西", work_order_id="WO1", trade="W1",
              finish="20261003120000", service="叠加包", valid="月包",
              product_id="WP1", product_name="工单产品", lh="去话",
              channel="热线人工", pos="一线坐席-10010", depart="10010",
              region="南方二中心", data_type="fwgd"):
    return [kf, tenant, work_order_id, trade, finish, service, "流量", valid,
            product_id, product_name, lh, channel, pos, depart, region, data_type]


def _org_row(kf="KF001", org="广西分公司->省中心->客户服务营销中心->热线运营组",
             finish="20261005120000", service="叠加包", valid="月包",
             product_id="OP1", product_name="线上直办产品", trade="O1"):
    return [kf, org, finish, service, valid, product_id, product_name, trade]


def _upload(headers, rows, pre_rows=None, filename="marketing.xlsx"):
    data = {"file": (filename, _make_xlsx(headers, rows, pre_rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    return client.post("/api/marketing/details/import", files=data)


def _set_overall_rules(rules):
    resp = client.put("/api/marketing/overall-row-rules", json={"rules": rules})
    assert resp.status_code == 200, resp.text
    return resp.json()


def _overall_row(report, label):
    return next(row for row in report["rows"] if row["label"] == label)


class TestMarketingImport:

    def setup_method(self):
        _clean()
        _seed_employees()

    def test_import_english_header_and_report(self):
        rows = [
            _en_row(trade="T1", product_id="P1", service="叠加包", valid="月包", lh="来话"),
            _en_row(trade="T2", product_id="P2", service="叠加包", valid="日包", lh="去话"),
            _en_row(kf="KF002", org="南方二中心>>通软-在线>>李智", trade="T3", product_id="P3",
                    service="套餐适配", valid="", lh="来话"),
            _en_row(kf="KF999", org="广西分公司>>省中心>>客户服务部>>客户关系运营中心",
                    trade="T4", product_id="P4", service="叠加包", valid="长期", lh="来话"),
        ]
        resp = _upload(EN_HEADERS, rows)
        assert resp.status_code == 200, resp.text
        assert resp.json()["count"] == 4

        _set_overall_rules([
            {"row_label": "线上直办-区域", "field": "org_name", "match_type": "contains", "pattern": "南方二中心", "priority": 1, "enabled": True},
            {"row_label": "线上直办-维系", "field": "org_name", "match_type": "contains", "pattern": "客户关系运营中心", "priority": 2, "enabled": True},
        ])
        report = client.get("/api/marketing/report", params={"year_month": "2026-10"}).json()
        assert report["stats"]["total"] == 4
        assert report["stats"]["overlay"] == 3
        assert report["stats"]["package_change"] == 1
        assert report["stats"]["inbound"] == 3
        assert report["stats"]["outbound"] == 1
        assert report["stats"]["people_count"] == 3
        assert set(report["options"]["groups"]) == {"未分组", "线上直办-区域", "线上直办-维系"}
        emp = next(item for item in report["employees"] if item["kf_id"] == "KF999")
        assert emp["name"] == ""
        assert emp["team"] == "未匹配班组"
        assert emp["group_name"] == "线上直办-维系"

    def test_import_chinese_header_and_same_date_overwrite(self):
        first = _upload(CN_HEADERS, [_cn_row(product_id="OLD", product_name="旧产品")])
        assert first.status_code == 200, first.text
        second = _upload(CN_HEADERS, [_cn_row(product_id="NEW", product_name="新产品")])
        assert second.status_code == 200, second.text
        details = client.get("/api/marketing/details", params={"start_date": "2026-10-02", "end_date": "2026-10-02"}).json()
        assert details["total"] == 1
        assert details["items"][0]["product_id"] == "NEW"

    def test_import_detects_header_after_title_rows(self):
        resp = _upload(CN_HEADERS, [_cn_row()], pre_rows=[["营销明细"], ["导出时间", "2026-10-20"]])
        assert resp.status_code == 200, resp.text
        assert resp.json()["count"] == 1

    def test_import_invalid_file(self):
        data = {"file": ("bad.xlsx", io.BytesIO(b"not excel"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        resp = client.post("/api/marketing/details/import", files=data)
        assert resp.status_code == 400

    def test_import_fwgd_source_and_same_date_does_not_overwrite_xszb(self):
        assert _upload(EN_HEADERS, [_en_row(trade="T1", product_id="P1", finish="20261004120000")]).status_code == 200
        resp = _upload(FWGD_HEADERS, [_fwgd_row(kf="KF001", trade="W1", product_id="WP1", finish="20261004120000")])
        assert resp.status_code == 200, resp.text
        assert resp.json()["data_type"] == "fwgd"
        assert resp.json()["source"] == "办理工单"

        details = client.get("/api/marketing/details", params={"start_date": "2026-10-04", "end_date": "2026-10-04"}).json()
        assert details["total"] == 2
        sources = {item["data_type"] for item in details["items"]}
        assert sources == {"xszb", "fwgd"}

        again = _upload(FWGD_HEADERS, [_fwgd_row(kf="KF001", trade="W2", product_id="WP2", finish="20261004120000")])
        assert again.status_code == 200, again.text
        details = client.get("/api/marketing/details", params={"start_date": "2026-10-04", "end_date": "2026-10-04"}).json()
        assert details["total"] == 2
        fwgd_items = [item for item in details["items"] if item["data_type"] == "fwgd"]
        assert len(fwgd_items) == 1
        assert fwgd_items[0]["trade_id"] == "W2"
        assert fwgd_items[0]["tenant"] == "广西"


class TestMarketingReport:

    def setup_method(self):
        _clean()
        _seed_employees()

    def test_business_type_valid_type_and_lh_filters(self):
        rows = [
            _en_row(trade="T1", product_id="P1", service="叠加包", valid="日包", lh="去话"),
            _en_row(trade="T2", product_id="P2", service="叠加包", valid="月包", lh="来话"),
            _en_row(trade="T3", product_id="P3", service="叠加包", valid="长期", lh="来话"),
            _en_row(trade="T4", product_id="P4", service="叠加包", valid="", lh="来话"),
            _en_row(trade="T5", product_id="P5", service="套餐适配", valid="", lh="来话"),
        ]
        assert _upload(EN_HEADERS, rows).status_code == 200
        report = client.get("/api/marketing/report", params={"start_date": "2026-10-01", "end_date": "2026-10-01"}).json()
        stats = report["stats"]
        assert stats["overlay"] == 4
        assert stats["overlay_day"] == 1
        assert stats["overlay_month"] == 1
        assert stats["overlay_long"] == 1
        assert stats["overlay_undefined"] == 1
        assert stats["package_change"] == 1

        filtered = client.get("/api/marketing/report", params={
            "start_date": "2026-10-01", "end_date": "2026-10-01",
            "business_type": "叠加包", "valid_type": "日包", "lh_qh": "去话",
        }).json()
        assert filtered["stats"]["total"] == 1
        assert filtered["employees"][0]["kf_id"] == "KF001"

    def test_overall_rule_priority_and_unclassified_group(self):
        rows = [
            _en_row(kf="KF001", org="A>>维系>>投诉中心", trade="T1", product_id="P1"),
            _en_row(kf="KF002", org="未知组织", trade="T2", product_id="P2"),
        ]
        assert _upload(EN_HEADERS, rows).status_code == 200
        _set_overall_rules([
            {"row_label": "线上直办-热线", "field": "org_name", "match_type": "contains", "pattern": "投诉中心", "priority": 1, "enabled": True},
            {"row_label": "线上直办-维系", "field": "org_name", "match_type": "contains", "pattern": "维系", "priority": 10, "enabled": True},
        ])
        report = client.get("/api/marketing/report", params={"year_month": "2026-10"}).json()
        groups = {item["group_name"]: item for item in report["groups"]}
        assert groups["线上直办-热线"]["total"] == 1
        assert groups["未分组"]["total"] == 1
        assert "线上直办-维系" not in groups
        rules = client.get("/api/marketing/overall-row-rules").json()
        assert len(rules["rules"]) == 2

    def test_regex_rule_is_shared_by_report_grouping(self):
        assert _upload(EN_HEADERS, [_en_row(org="广西分公司>>热线运营组", trade="T1", product_id="P1")]).status_code == 200
        _set_overall_rules([
            {"row_label": "线上直办-热线", "field": "org_name", "match_type": "regex", "pattern": "热线|营销", "priority": 1, "enabled": True},
        ])
        report = client.get("/api/marketing/report", params={"year_month": "2026-10"}).json()
        assert report["employees"][0]["group_name"] == "线上直办-热线"
        assert "线上直办-热线" in report["options"]["groups"]

    def test_overall_report_structure_subtotal_and_unclassified(self):
        rows = [
            _en_row(kf="KF001", org="广西分公司>>热线运营组", trade="T1", product_id="P1", service="叠加包", valid="月包"),
            _en_row(kf="KF002", org="南方二中心>>通软-在线", trade="T2", product_id="P2", service="套餐适配", valid=""),
            _en_row(kf="KF001", org="未知组织", trade="T3", product_id="P3", service="叠加包", valid="日包"),
        ]
        assert _upload(EN_HEADERS, rows).status_code == 200
        _set_overall_rules([
            {"row_label": "线上直办-热线", "field": "org_name", "match_type": "contains", "pattern": "热线运营组", "priority": 1, "enabled": True},
            {"row_label": "线上直办-区域", "field": "org_name", "match_type": "contains", "pattern": "南方二中心", "priority": 2, "enabled": True},
        ])
        report = client.get("/api/marketing/overall-report", params={"year_month": "2026-10"}).json()
        assert len(report["rows"]) == 13
        hotline = _overall_row(report, "线上直办-热线")
        area = _overall_row(report, "线上直办-区域")
        sub1 = next(row for row in report["rows"] if row["section_no"] == 1 and row["is_subtotal"])
        assert hotline["counts"]["overlay"] == 1
        assert area["counts"]["package_change"] == 1
        assert sub1["counts"]["overlay"] == 1
        assert sub1["counts"]["package_change"] == 1
        assert sub1["total"] == 2
        assert report["stats"]["grand_total"] == 2
        assert report["stats"]["unclassified_rows"] == 1
        assert report["total_row"]["total"] == 2
        assert hotline["rowspan"] == 6
        assert report["rows"][1]["rowspan"] == 0

    def test_overall_source_data_ranges_global_and_query(self):
        online_rows = [
            _en_row(trade="ON1", product_id="P1", finish="20260930120000"),
            _en_row(trade="ON2", product_id="P2", finish="20261001120000"),
            _en_row(trade="ON3", product_id="P3", finish="20261020120000"),
            _en_row(trade="ON4", product_id="P4", finish="20261101120000"),
        ]
        assert _upload(EN_HEADERS, online_rows).status_code == 200
        work_order_rows = [
            _fwgd_row(trade="WO1", product_id="W1", finish="20261002120000"),
            _fwgd_row(trade="WO2", product_id="W2", finish="20261022120000"),
        ]
        assert _upload(FWGD_HEADERS, work_order_rows).status_code == 200

        report = client.get("/api/marketing/overall-report", params={
            "start_date": "2026-10-01",
            "end_date": "2026-10-31",
        }).json()
        assert report["stats"]["start_date"] == "2026-10-01"
        assert report["stats"]["end_date"] == "2026-10-31"
        assert report["data_ranges"]["database"] == {
            "xszb": {"start_date": "2026-09-30", "end_date": "2026-11-01"},
            "fwgd": {"start_date": "2026-10-02", "end_date": "2026-10-22"},
        }
        assert report["data_ranges"]["query"] == {
            "xszb": {"start_date": "2026-10-01", "end_date": "2026-10-20"},
            "fwgd": {"start_date": "2026-10-02", "end_date": "2026-10-22"},
        }

        monthly = client.get("/api/marketing/overall-report", params={"year_month": "2026-10"}).json()
        assert monthly["data_ranges"]["database"] == report["data_ranges"]["database"]

        empty = client.get("/api/marketing/overall-report", params={
            "start_date": "2026-10-10",
            "end_date": "2026-10-15",
        }).json()
        assert empty["data_ranges"]["database"] == report["data_ranges"]["database"]
        assert empty["data_ranges"]["query"] == {
            "xszb": {"start_date": None, "end_date": None},
            "fwgd": {"start_date": None, "end_date": None},
        }

    def test_overall_export_has_style_merge_and_formulas(self):
        assert _upload(EN_HEADERS, [_en_row(org="热线运营组")]).status_code == 200
        _set_overall_rules([
            {"row_label": "线上直办-热线", "field": "org_name", "match_type": "contains", "pattern": "热线运营组", "priority": 1, "enabled": True},
        ])
        resp = client.get("/api/marketing/overall-report/export", params={"year_month": "2026-10"})
        assert resp.status_code == 200
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb["总体营销"]
        assert ws["A1"].value == "序号"
        assert ws["C1"].font.color.rgb == "FFFF0000"
        assert ws["A2"].value == 1
        assert any(str(rng) == "A2:A7" for rng in ws.merged_cells.ranges)
        assert ws["H2"].value == "=SUM(C2:G2)"
        assert str(ws["I2"].value).startswith("=IFERROR")
        assert ws["A15"].value == "合计"
        assert ws["B2"].value == "线上直办-热线"

    def test_delete_details_by_date_and_batch(self):
        resp = _upload(EN_HEADERS, [_en_row()])
        batch = resp.json()["batch"]
        details = client.get("/api/marketing/details").json()
        assert details["total"] == 1
        deleted = client.delete(f"/api/marketing/details/import/{batch}")
        assert deleted.status_code == 200
        assert deleted.json()["count"] == 1
        assert client.get("/api/marketing/details").json()["total"] == 0


class TestMarketingWorkOrders:

    def setup_method(self):
        _clean()
        _seed_employees()

    def test_fwgd_uses_xszb_organization_before_tenant_fallback(self):
        assert _upload(ORG_HEADERS, [_org_row(kf="KF001", trade="O1", product_id="OP1")]).status_code == 200
        _set_overall_rules([
            {"row_label": "营销工单-区域", "field": "tenant", "match_type": "contains", "pattern": "南方二中心", "priority": 0, "enabled": True},
            {"row_label": "线上直办-热线", "field": "org_name", "match_type": "contains", "pattern": "热线运营组", "priority": 99, "enabled": True},
        ])
        resp = _upload(FWGD_HEADERS, [
            _fwgd_row(kf="KF001", trade="W1", product_id="WP1", finish="20261005120000"),
            _fwgd_row(kf="KF999", tenant="南方二中心", trade="W2", product_id="WP2", finish="20261005120000", lh="来话"),
        ])
        assert resp.status_code == 200, resp.text
        assert resp.json()["data_type"] == "fwgd"

        report = client.get("/api/marketing/overall-report", params={"year_month": "2026-10"}).json()
        assert len(report["rows"]) == 13
        assert all(row["label"] != "办理工单" for row in report["rows"])
        assert _overall_row(report, "线上直办-热线")["counts"]["overlay"] == 1
        assert _overall_row(report, "营销工单-热线")["counts"]["overlay"] == 1
        assert _overall_row(report, "营销工单-区域")["counts"]["overlay"] == 1
        assert report["stats"]["grand_total"] == 3
        assert report["stats"]["unclassified_rows"] == 0

    def test_fwgd_unclassified_not_counted_in_overall(self):
        assert _upload(FWGD_HEADERS, [_fwgd_row(kf="KF999", tenant="未知租户")]).status_code == 200
        report = client.get("/api/marketing/overall-report", params={"year_month": "2026-10"}).json()
        assert report["stats"]["grand_total"] == 0
        assert report["stats"]["unclassified_rows"] == 1
        assert all(row["total"] == 0 for row in report["rows"])

    def test_marketing_report_source_filter(self):
        assert _upload(ORG_HEADERS, [_org_row(kf="KF001", trade="O1", product_id="OP1")]).status_code == 200
        assert _upload(FWGD_HEADERS, [_fwgd_row(kf="KF001", trade="W1", product_id="WP1", finish="20261005120000")]).status_code == 200

        all_report = client.get("/api/marketing/report", params={"year_month": "2026-10"}).json()
        assert all_report["stats"]["total"] == 2
        assert {item["value"] for item in all_report["options"]["sources"]} == {"xszb", "fwgd"}

        fwgd_report = client.get("/api/marketing/report", params={"year_month": "2026-10", "data_type": "fwgd"}).json()
        assert fwgd_report["stats"]["total"] == 1
        assert fwgd_report["stats"]["outbound"] == 1

        xszb_report = client.get("/api/marketing/report", params={"year_month": "2026-10", "data_type": "xszb"}).json()
        assert xszb_report["stats"]["total"] == 1
        assert xszb_report["stats"]["inbound"] == 0

        details = client.get("/api/marketing/details", params={"start_date": "2026-10-01", "end_date": "2026-10-31", "data_type": "fwgd"}).json()
        assert details["total"] == 1
        assert details["items"][0]["source"] == "办理工单"
        assert details["items"][0]["tenant"] == "广西"


class TestKeyProductStatistics:

    def setup_method(self):
        _clean()
        _seed_employees()

    def _seed_key_products(self):
        resp = client.put("/api/marketing/key-products", json={
            "products": [
                {"business_name": "业务A", "product_id": "P1", "product_name": "产品A"},
                {"business_name": "业务B", "product_id": "P2", "product_name": "产品B"},
                {"business_name": "业务C", "product_id": "P3", "product_name": "产品C"},
            ],
        })
        assert resp.status_code == 200, resp.text
        return resp.json()

    def _seed_data(self):
        online_rows = [
            _en_row(kf="KF001", org="广西分公司>>热线运营组", trade="T1", product_id="P1", product_name="产品A"),
            _en_row(kf="KF001", org="广西分公司>>热线运营组", trade="T2", product_id="P2", product_name="产品B"),
            _en_row(kf="KF002", org="南方二中心>>通软-在线", trade="T3", product_id="P1", product_name="产品A"),
            _en_row(kf="KF999", org="广西分公司>>热线运营组", trade="T4", product_id="P1", product_name="产品A"),
        ]
        assert _upload(EN_HEADERS, online_rows).status_code == 200
        assert _upload(FWGD_HEADERS, [
            _fwgd_row(kf="KF002", tenant="广西", trade="W1", product_id="P3", product_name="产品C", finish="20261005120000"),
        ]).status_code == 200
        _set_overall_rules([
            {"row_label": "线上直办-热线", "field": "org_name", "match_type": "contains", "pattern": "热线运营组", "priority": 1, "enabled": True},
            {"row_label": "线上直办-区域", "field": "org_name", "match_type": "contains", "pattern": "南方二中心", "priority": 2, "enabled": True},
        ])

    def _report(self, source):
        return client.get("/api/marketing/overall-report", params={
            "year_month": "2026-10",
            "product_source": source,
        }).json()

    def test_key_product_config_crud_and_validation(self):
        resp = client.put("/api/marketing/key-products", json={
            "products": [
                {"business_name": "业务B", "product_id": "P2", "product_name": "产品B"},
                {"business_name": "业务A", "product_name": "产品A"},
            ],
        })
        assert resp.status_code == 200, resp.text
        products = resp.json()["products"]
        assert [item["business_name"] for item in products] == ["业务B", "业务A"]
        assert [item["sort"] for item in products] == [0, 1]
        assert products[0]["id"]

        fetched = client.get("/api/marketing/key-products").json()["products"]
        assert [item["business_name"] for item in fetched] == ["业务B", "业务A"]

        assert client.put("/api/marketing/key-products", json={
            "products": [{"business_name": "", "product_id": "P1"}],
        }).status_code == 400
        assert client.put("/api/marketing/key-products", json={
            "products": [{"business_name": "无产品"}],
        }).status_code == 400
        assert client.put("/api/marketing/key-products", json={
            "products": [
                {"business_name": "重复1", "product_id": "P1"},
                {"business_name": "重复2", "product_id": "p1"},
            ],
        }).status_code == 400

    def test_product_options_returns_distinct_products(self):
        self._seed_data()
        items = client.get("/api/marketing/product-options").json()["items"]
        lookup = {(item["product_id"], item["product_name"]): item["count"] for item in items}
        assert lookup[("P1", "产品A")] == 3
        assert lookup[("P2", "产品B")] == 1
        assert lookup[("P3", "产品C")] == 1

        filtered = client.get("/api/marketing/product-options", params={"keyword": "产品B"}).json()["items"]
        assert len(filtered) == 1
        assert filtered[0]["product_id"] == "P2"

    def test_product_table_defaults_to_online_source(self):
        self._seed_key_products()
        self._seed_data()
        report = self._report("xszb")
        assert report["key_product_source"] == "xszb"
        assert report["key_product_source_label"] == "线上直办"
        assert [col["label"] for col in report["key_product_columns"]] == ["业务A", "业务B", "业务C"]
        rows = report["key_product_rows"]
        assert [row["work_group"] for row in rows if row["rowspan"]] == ["热线", "区域"]
        hotline_team = next(row for row in rows if row["work_group"] == "热线" and row["team"] == "热线一组")
        assert hotline_team["counts"] == {"key_product_0": 1, "key_product_1": 1, "key_product_2": 0}
        assert hotline_team["total"] == 2
        assert hotline_team["rowspan"] == 3
        unassigned = next(row for row in rows if row["team"] == "未分班组")
        assert unassigned["work_group"] == "热线"
        assert unassigned["counts"]["key_product_0"] == 1
        hotline_subtotal = next(row for row in rows if row["work_group"] == "热线" and row.get("is_subtotal"))
        assert hotline_subtotal["team"] == "小计"
        assert hotline_subtotal["counts"] == {"key_product_0": 2, "key_product_1": 1, "key_product_2": 0}
        assert hotline_subtotal["total"] == 3
        assert hotline_subtotal["rowspan"] == 0
        area_team = next(row for row in rows if row["work_group"] == "区域")
        assert area_team["team"] == "区域二组"
        assert area_team["counts"]["key_product_0"] == 1
        assert not any(row["work_group"] == "区域" and row.get("is_subtotal") for row in rows)
        assert all(row["counts"].get("key_product_2", 0) == 0 for row in rows)
        assert report["key_product_total_row"] == {
            "counts": {"key_product_0": 3, "key_product_1": 1, "key_product_2": 0},
            "total": 4,
        }

    def test_product_table_supports_work_order_and_total_sources(self):
        self._seed_key_products()
        self._seed_data()
        work_order = self._report("fwgd")
        work_order_rows = work_order["key_product_rows"]
        assert len(work_order_rows) == 1
        assert work_order_rows[0]["work_group"] == "区域"
        assert work_order_rows[0]["team"] == "区域二组"
        assert work_order_rows[0]["counts"] == {"key_product_0": 0, "key_product_1": 0, "key_product_2": 1}
        assert work_order["key_product_total_row"]["total"] == 1

        total = self._report("all")
        assert total["key_product_source_label"] == "合计"
        assert total["key_product_total_row"] == {
            "counts": {"key_product_0": 3, "key_product_1": 1, "key_product_2": 1},
            "total": 5,
        }
        assert client.get("/api/marketing/overall-report", params={
            "year_month": "2026-10", "product_source": "bad",
        }).status_code == 400

    def test_product_table_export_contains_new_table(self):
        self._seed_key_products()
        self._seed_data()
        resp = client.get("/api/marketing/overall-report/export", params={
            "year_month": "2026-10", "product_source": "all",
        })
        assert resp.status_code == 200
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb["总体营销"]
        values = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None]
        assert "重点业务产品统计（合计）" in values
        assert "工作组" in values
        assert "业务A" in values
        assert "未分班组" in values
        assert "小计" in values
        all_rows = list(ws.iter_rows())
        title_index = next(index for index, row in enumerate(all_rows) if row[0].value == "重点业务产品统计（合计）")
        subtotal_row = next(row for row in all_rows[title_index:] if row[1].value == "小计")
        assert [cell.value for cell in subtotal_row[2:5]] == [2, 1, 0]


class TestMarketingRulesAndExports:

    def setup_method(self):
        _clean()
        _seed_employees()

    def test_overall_rule_validation(self):
        resp = client.put("/api/marketing/overall-row-rules", json={
            "rules": [{"row_label": "线上直办-热线", "field": "org_name", "match_type": "unknown", "pattern": "x", "priority": 1}],
        })
        assert resp.status_code == 400
        resp = client.put("/api/marketing/overall-row-rules", json={
            "rules": [{"row_label": "不存在", "field": "org_name", "match_type": "contains", "pattern": "x"}],
        })
        assert resp.status_code == 400
        resp = client.put("/api/marketing/overall-row-rules", json={
            "rules": [{"row_label": "线上直办-热线", "field": "unknown", "match_type": "contains", "pattern": "x"}],
        })
        assert resp.status_code == 400

    def test_permissions_required_for_mutation_and_export(self):
        original = app.dependency_overrides[get_current_user]
        app.dependency_overrides[get_current_user] = lambda: {
            "id": 2, "username": "limited", "role": "user", "is_system": False, "permissions": "{}"
        }
        try:
            assert client.get("/api/marketing/report").status_code == 403
            assert client.get("/api/marketing/overall-report").status_code == 403
            assert client.get("/api/marketing/report/export").status_code == 403
            assert client.get("/api/marketing/overall-report/export").status_code == 403
            assert client.put("/api/marketing/overall-row-rules", json={"rules": []}).status_code == 403
            data = {"file": ("x.xlsx", _make_xlsx(EN_HEADERS, [_en_row()]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
            assert client.post("/api/marketing/details/import", files=data).status_code == 403
        finally:
            app.dependency_overrides[get_current_user] = original

    def test_report_csv_export(self):
        rows = [
            _en_row(kf="KF001", trade="T1", product_id="P1"),
            _en_row(kf="KF002", org="未知组织", trade="T2", product_id="P2"),
        ]
        assert _upload(EN_HEADERS, rows).status_code == 200
        _set_overall_rules([])
        resp = client.get("/api/marketing/report/export", params={"year_month": "2026-10", "dimension": "group"})
        assert resp.status_code == 200
        assert "text/csv" in resp.headers["content-type"]
        text = resp.content.decode("utf-8-sig")
        assert "行分类" in text
        assert "未分组" in text
        assert "产品办理量" in text

    def test_report_team_includes_leaders_and_csv_export(self):
        _seed_leader("L001", "王组长", "热线一组")
        _seed_leader("L002", "赵组长", "热线一组")
        _seed_leader("L003", "离职组长", "热线一组", status="离职")
        assert _upload(EN_HEADERS, [
            _en_row(kf="KF001", trade="T1", product_id="P1"),
            _en_row(kf="KF999", trade="T2", product_id="P2"),
        ]).status_code == 200

        report = client.get("/api/marketing/report", params={"year_month": "2026-10"}).json()
        team = next(item for item in report["teams"] if item["team"] == "热线一组")
        assert team["leader"] == "王组长、赵组长"
        unmatched = next(item for item in report["teams"] if item["team"] == "未匹配班组")
        assert unmatched["leader"] == ""

        resp = client.get("/api/marketing/report/export", params={"year_month": "2026-10", "dimension": "team"})
        assert resp.status_code == 200
        text = resp.content.decode("utf-8-sig")
        assert text.splitlines()[0].startswith("班组,组长,人数,")
        assert "热线一组,王组长、赵组长," in text

    def test_import_deduplicates_same_trade_product(self):
        rows = [
            _en_row(trade="T1", product_id="P1", product_name="产品A"),
            _en_row(trade="T1", product_id="P1", product_name="产品A"),
        ]
        resp = _upload(EN_HEADERS, rows)
        assert resp.status_code == 200, resp.text
        assert resp.json()["count"] == 1
        assert resp.json()["skipped"] == 1

    def test_overall_export_zero_data_uses_green_ratio(self):
        resp = client.get("/api/marketing/overall-report/export", params={"year_month": "2026-10"})
        assert resp.status_code == 200
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb["总体营销"]
        assert ws["I2"].font.color.rgb == "FF00B050"
        assert ws["I2"].fill.start_color.rgb == "FFC6EFCE"
        assert ws["H15"].value == "=SUM(C15:G15)"
